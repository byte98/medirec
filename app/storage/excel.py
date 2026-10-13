# MediRec - simple web application to save some basic medical records.
# Copyright (C) 2026 Jiri Skoda <developer@skodaj.cz>
#
# This program is free software; you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation; version 2 of the License.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License along
# with this program; if not, write to the Free Software Foundation, Inc.,
# 51 Franklin Street, Fifth Floor, Boston, MA 02110-1301 USA.

"""Crash-safe and multi-process-safe XLSX storage."""

from __future__ import annotations

import os
import copy
import stat
import shutil
import tempfile
from io import BytesIO
import unicodedata
import uuid
from contextlib import contextmanager
from pathlib import Path
from threading import Lock
from typing import Callable, Iterator, TypeVar

import fcntl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import GradientFill, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from ..models.data import ClientCreate, ClientDetail, ClientSummary, HealthInsuranceCompany, HealthMeasurement, HealthMeasurementCreate
from ..health.indicators import IndicatorCalculator

T = TypeVar("T")


class WorkbookStorage:
    """Repository for the MediRec XLSX workbook.

    Every write obtains an exclusive inter-process lock, modifies a freshly
    loaded workbook, writes it to a temporary file in the workbook directory,
    validates that temporary workbook can be opened, fsyncs it and atomically
    replaces the original file. This prevents a partially written XLSX from
    becoming the live database after a process crash or power loss.
    """

    HEALTH_INSURANCE_SHEET = "HealthInsuranceCompanies"
    TEMPLATE_SHEET = "_template"
    FIRST_HIC_ROW = 5
    FIRST_DATA_ROW = 16

    def __init__(self, workbook_path: Path):
        """Create storage for *workbook_path*.

        :param workbook_path: XLSX database path.
        """
        self._workbook_path = workbook_path
        self._lock_path = workbook_path.with_name(workbook_path.name + ".lock")
        self._thread_lock = Lock()

    @property
    def workbook_path(self) -> Path:
        """Return the configured workbook path."""
        reti = self._workbook_path
        return reti

    def initialize(self) -> None:
        """Validate that the workbook has the required sheets."""
        if not self._workbook_path.exists():
            raise FileNotFoundError(f"Workbook not found: {self._workbook_path}")
        with self._locked(fcntl.LOCK_SH):
            workbook = load_workbook(self._workbook_path, data_only=False)
            try:
                missing = [name for name in (self.HEALTH_INSURANCE_SHEET, self.TEMPLATE_SHEET) if name not in workbook.sheetnames]
                if missing:
                    raise ValueError("Workbook is missing required sheets: " + ", ".join(missing))
            finally:
                workbook.close()

    @contextmanager
    def _locked(self, mode: int) -> Iterator[None]:
        """Acquire both an in-process and inter-process workbook lock."""
        with self._thread_lock:
            self._lock_path.parent.mkdir(parents=True, exist_ok=True)
            with self._lock_path.open("a+b") as lock_file:
                fcntl.flock(lock_file.fileno(), mode)
                try:
                    yield
                finally:
                    fcntl.flock(lock_file.fileno(), fcntl.LOCK_UN)

    def _read(self, callback: Callable[[Workbook], T]) -> T:
        """Execute a read-only callback against a workbook snapshot."""
        with self._locked(fcntl.LOCK_SH):
            workbook = load_workbook(self._workbook_path, data_only=False)
            try:
                result = callback(workbook)
            finally:
                workbook.close()
        return result

    def _write(self, callback: Callable[[Workbook], T]) -> T:
        """Execute an atomic workbook transaction."""
        with self._locked(fcntl.LOCK_EX):
            workbook = load_workbook(self._workbook_path, data_only=False)
            try:
                result = callback(workbook)
                self._atomic_save(workbook)
            finally:
                workbook.close()
        return result

    def _atomic_save(self, workbook: Workbook) -> None:
        """Save and atomically publish a workbook."""
        self._workbook_path.parent.mkdir(parents=True, exist_ok=True)
        original_mode = stat.S_IMODE(self._workbook_path.stat().st_mode)
        file_descriptor, temporary_name = tempfile.mkstemp(prefix="." + self._workbook_path.name + ".", suffix=".xlsx.tmp", dir=self._workbook_path.parent)
        os.close(file_descriptor)
        temporary_path = Path(temporary_name)
        try:
            workbook.save(temporary_path)
            os.chmod(temporary_path, original_mode)
            with temporary_path.open("rb") as temporary_file:
                os.fsync(temporary_file.fileno())
            validation = load_workbook(BytesIO(temporary_path.read_bytes()), read_only=True, data_only=False)
            validation.close()
            os.replace(temporary_path, self._workbook_path)
            directory_descriptor = os.open(self._workbook_path.parent, os.O_RDONLY)
            try:
                os.fsync(directory_descriptor)
            finally:
                os.close(directory_descriptor)
        finally:
            if temporary_path.exists():
                temporary_path.unlink()

    @staticmethod
    def _ascii_sheet_part(value: str) -> str:
        """Convert a name to the nearest safe ASCII sheet-name component."""
        replacements = {"Ł": "L", "ł": "l", "Đ": "D", "đ": "d", "Ð": "D", "ð": "d", "Þ": "Th", "þ": "th", "Æ": "AE", "æ": "ae", "Œ": "OE", "œ": "oe", "Ø": "O", "ø": "o"}
        translated = "".join(replacements.get(character, character) for character in value)
        normalized = unicodedata.normalize("NFKD", translated)
        ascii_value = normalized.encode("ascii", "ignore").decode("ascii")
        safe = "".join(character if character not in "[]:*?/\\" else "_" for character in ascii_value)
        safe = "_".join(safe.split())
        reti = safe or "Unknown"
        return reti

    @classmethod
    def build_sheet_name(cls, client: ClientCreate) -> str:
        """Build the required ASCII sheet name from surname, name and personal ID.

        Excel limits worksheet names to 31 characters. The personal ID is kept
        intact whenever possible, while surname and name are shortened only
        when the complete requested name would exceed that limit.

        :param client: Client whose personal data are used.
        :returns: Valid Excel worksheet name.
        """
        surname = cls._ascii_sheet_part(client.surname)
        name = cls._ascii_sheet_part(client.name)
        personal_id = cls._ascii_sheet_part(client.personal_id)
        base = f"{surname}_{name}_{personal_id}"
        if len(base) <= 31:
            reti = base
        else:
            suffix = f"_{personal_id}"
            available = max(1, 31 - len(suffix) - 1)
            combined = f"{surname}_{name}"
            prefix = combined[:available].rstrip("_")
            reti = f"{prefix}{suffix}"[:31]
        return reti

    @classmethod
    def _client_from_sheet(cls, sheet: Worksheet) -> ClientSummary:
        """Read personal data from a client sheet."""
        value = lambda coordinate: sheet[coordinate].value
        date_value = value("C6")
        if hasattr(date_value, "date"):
            date_value = date_value.date()
        client = ClientSummary(
            surname=str(value("C2") or ""),
            name=str(value("C4") or ""),
            identifier=str(value("N2") or ""),
            date_of_birth=date_value,
            personal_id=str(value("C8") or ""),
            sex=str(value("J8") or ""),
            health_insurance_abbreviation=str(value("E10") or ""),
            sheet_name=sheet.title,
        )
        return client

    @staticmethod
    def _value(value: object) -> object:
        """Convert empty Excel values to None."""
        reti = value
        if value == "":
            reti = None
        return reti

    def list_clients(self) -> list[ClientSummary]:
        """Return all client sheets."""
        def read(workbook: Workbook) -> list[ClientSummary]:
            clients = []
            for name in workbook.sheetnames:
                if name not in (self.HEALTH_INSURANCE_SHEET, self.TEMPLATE_SHEET):
                    clients.append(self._client_from_sheet(workbook[name]))
            clients.sort(key=lambda item: (item.surname.casefold(), item.name.casefold()))
            reti = clients
            return reti
        reti = self._read(read)
        return reti

    def list_health_insurance_companies(self) -> list[HealthInsuranceCompany]:
        """Return all health insurance companies."""
        def read(workbook: Workbook) -> list[HealthInsuranceCompany]:
            sheet = workbook[self.HEALTH_INSURANCE_SHEET]
            companies = []
            row = self.FIRST_HIC_ROW
            while row <= sheet.max_row:
                abbreviation = sheet.cell(row, 1).value
                name = sheet.cell(row, 2).value
                code = sheet.cell(row, 3).value
                notes = sheet.cell(row, 4).value
                if any(value not in (None, "") for value in (abbreviation, name, code, notes)):
                    companies.append(HealthInsuranceCompany(abbreviation=str(abbreviation or ""), name=str(name or ""), code=str(code or ""), notes=str(notes or "")))
                row += 1
            reti = companies
            return reti
        reti = self._read(read)
        return reti

    def add_health_insurance_company(self, company: HealthInsuranceCompany) -> None:
        """Append a health insurance company to the first free row."""
        def write(workbook: Workbook) -> None:
            sheet = workbook[self.HEALTH_INSURANCE_SHEET]
            row = self.FIRST_HIC_ROW
            while any(sheet.cell(row, column).value not in (None, "") for column in range(1, 5)):
                row += 1
            existing = {str(sheet.cell(item, 1).value).casefold() for item in range(self.FIRST_HIC_ROW, sheet.max_row + 1) if sheet.cell(item, 1).value not in (None, "")}
            if company.abbreviation.casefold() in existing:
                raise ValueError("A health insurance company with this abbreviation already exists.")
            sheet.cell(row, 1).value = company.abbreviation
            sheet.cell(row, 2).value = company.name
            sheet.cell(row, 3).value = company.code
            sheet.cell(row, 4).value = company.notes
        self._write(write)

    def create_client(self, client: ClientCreate) -> ClientSummary:
        """Create a client sheet by copying the template and generating its identifier."""
        def write(workbook: Workbook) -> ClientSummary:
            insurance_codes = set()
            for row in range(self.FIRST_HIC_ROW, workbook[self.HEALTH_INSURANCE_SHEET].max_row + 1):
                insurance_codes.add(str(workbook[self.HEALTH_INSURANCE_SHEET].cell(row, 1).value or "").casefold())
            if client.health_insurance_abbreviation.casefold() not in insurance_codes:
                raise ValueError("Unknown health insurance company abbreviation.")

            identifier = uuid.uuid4().hex
            while any(
                name not in (self.HEALTH_INSURANCE_SHEET, self.TEMPLATE_SHEET)
                and str(workbook[name]["N2"].value or "") == identifier
                for name in workbook.sheetnames
            ):
                identifier = uuid.uuid4().hex

            sheet_name = self.build_sheet_name(client)
            if sheet_name in workbook.sheetnames:
                raise ValueError("The generated sheet name is already in use.")
            template = workbook[self.TEMPLATE_SHEET]
            sheet = workbook.copy_worksheet(template)
            sheet.title = sheet_name
            sheet["C2"] = client.surname
            sheet["C4"] = client.name
            sheet["N2"] = identifier
            sheet["C6"] = client.date_of_birth
            sheet["C8"] = client.personal_id
            sheet["J8"] = client.sex.upper()
            sheet["E10"] = client.health_insurance_abbreviation
            reti = ClientSummary(**client.model_dump(), identifier=identifier, sheet_name=sheet_name)
            return reti
        reti = self._write(write)
        return reti

    def delete_client(self, identifier: str) -> None:
        """Delete the client sheet identified by the primary identifier."""
        def write(workbook: Workbook) -> None:
            target = None
            for name in workbook.sheetnames:
                if name not in (self.HEALTH_INSURANCE_SHEET, self.TEMPLATE_SHEET):
                    if str(workbook[name]["N2"].value or "") == identifier:
                        target = name
            if target is None:
                raise KeyError("Client not found.")
            workbook.remove(workbook[target])
        self._write(write)

    def get_client(self, identifier: str) -> ClientDetail:
        """Return a client and all stored health measurements."""
        def read(workbook: Workbook) -> ClientDetail:
            target = None
            for name in workbook.sheetnames:
                if name not in (self.HEALTH_INSURANCE_SHEET, self.TEMPLATE_SHEET):
                    if str(workbook[name]["N2"].value or "") == identifier:
                        target = workbook[name]
            if target is None:
                raise KeyError("Client not found.")
            summary = self._client_from_sheet(target)
            measurements = self._read_measurements(target)
            reti = ClientDetail(**summary.model_dump(), measurements=measurements)
            return reti
        reti = self._read(read)
        return reti

    def _read_measurements(self, sheet: Worksheet) -> list[HealthMeasurement]:
        """Read health measurements from a client sheet."""
        rows = []
        previous_weight = None
        for row in range(self.FIRST_DATA_ROW, sheet.max_row + 1):
            values = [sheet.cell(row, column).value for column in range(1, 17)]
            if all(value in (None, "") for value in values):
                continue
            measurement_date = values[0]
            measurement_time = values[1]
            if hasattr(measurement_date, "date"):
                measurement_date = measurement_date.date()
            if hasattr(measurement_time, "time"):
                measurement_time = measurement_time.time()
            systolic = self._value(values[3])
            diastolic = self._value(values[5])
            pulse = self._value(values[7])
            oxygen = self._value(values[9])
            height = self._value(values[10])
            weight = self._value(values[12])
            bmi = self._value(values[14])
            item = HealthMeasurement(
                date=measurement_date,
                time=measurement_time,
                systolic=systolic,
                diastolic=diastolic,
                pulse=pulse,
                oxygen_saturation=oxygen,
                height=height,
                weight=weight,
                bmi=bmi,
                notes=str(values[15] or ""),
                systolic_status=IndicatorCalculator.systolic(systolic),
                diastolic_status=IndicatorCalculator.diastolic(diastolic),
                pulse_status=IndicatorCalculator.pulse(pulse),
                oxygen_status=IndicatorCalculator.oxygen(oxygen),
                weight_status=IndicatorCalculator.weight(weight, previous_weight),
                bmi_status=IndicatorCalculator.bmi(bmi),
            )
            rows.append(item)
            if weight is not None:
                previous_weight = weight
        reti = rows
        return reti


    @staticmethod
    def _set_indicator(cell: object, status: object) -> None:
        """Set a narrow colored rectangle in the dedicated indicator cell."""
        colors = {
            "HEALTHY": ("66BB6A", "2E7D32"),
            "CAUTION": ("F4D35E", "C49A00"),
            "UNHEALTHY": ("EF6C6C", "B71C1C"),
            "UNKNOWN": ("BDBDBD", "616161"),
        }
        status_name = getattr(status, "value", str(status))
        cell.value = None
        if status is None or status_name == "UNKNOWN":
            cell.fill = PatternFill(fill_type=None)
        else:
            start, end = colors.get(status_name, colors["UNKNOWN"])
            cell.fill = GradientFill(type="linear", degree=0, stop=[start, end])
        alignment = copy.copy(cell.alignment)
        alignment.horizontal = "center"
        alignment.vertical = "center"
        cell.alignment = alignment

    def add_measurement(self, identifier: str, measurement: HealthMeasurementCreate) -> HealthMeasurement:
        """Append a measurement to a client sheet."""
        def write(workbook: Workbook) -> HealthMeasurement:
            target = None
            for name in workbook.sheetnames:
                if name not in (self.HEALTH_INSURANCE_SHEET, self.TEMPLATE_SHEET):
                    if str(workbook[name]["N2"].value or "") == identifier:
                        target = workbook[name]
            if target is None:
                raise KeyError("Client not found.")
            previous_weight = None
            row = self.FIRST_DATA_ROW
            while any(target.cell(row, column).value not in (None, "") for column in range(1, 17)):
                if target.cell(row, 13).value not in (None, ""):
                    previous_weight = float(target.cell(row, 13).value)
                row += 1
            bmi = measurement.bmi
            if bmi is None and measurement.height and measurement.height > 0 and measurement.weight is not None:
                bmi = measurement.weight / ((measurement.height / 100) ** 2)
            for column in (3, 5, 7, 9, 12, 14):
                target.column_dimensions[target.cell(row, column).column_letter].width = 2.5
            target.cell(row, 1).value = measurement.date
            target.cell(row, 1).number_format = "dd.mm.yyyy"
            target.cell(row, 2).value = measurement.time
            target.cell(row, 2).number_format = "HH:mm"
            self._set_indicator(target.cell(row, 3), IndicatorCalculator.systolic(measurement.systolic))
            target.cell(row, 4).value = measurement.systolic
            self._set_indicator(target.cell(row, 5), IndicatorCalculator.diastolic(measurement.diastolic))
            target.cell(row, 6).value = measurement.diastolic
            self._set_indicator(target.cell(row, 7), IndicatorCalculator.pulse(measurement.pulse))
            target.cell(row, 8).value = measurement.pulse
            self._set_indicator(target.cell(row, 9), IndicatorCalculator.oxygen(measurement.oxygen_saturation))
            target.cell(row, 10).value = measurement.oxygen_saturation
            target.cell(row, 11).value = measurement.height
            self._set_indicator(target.cell(row, 12), IndicatorCalculator.weight(measurement.weight, previous_weight))
            target.cell(row, 13).value = measurement.weight
            self._set_indicator(target.cell(row, 14), IndicatorCalculator.bmi(bmi))
            target.cell(row, 15).value = bmi
            target.cell(row, 16).value = measurement.notes
            statuses = dict(
                systolic_status=IndicatorCalculator.systolic(measurement.systolic),
                diastolic_status=IndicatorCalculator.diastolic(measurement.diastolic),
                pulse_status=IndicatorCalculator.pulse(measurement.pulse),
                oxygen_status=IndicatorCalculator.oxygen(measurement.oxygen_saturation),
                weight_status=IndicatorCalculator.weight(measurement.weight, previous_weight),
                bmi_status=IndicatorCalculator.bmi(bmi),
            )
            reti = HealthMeasurement(**measurement.model_dump(exclude={"bmi"}), bmi=bmi, **statuses)
            return reti
        reti = self._write(write)
        return reti
