# MediRec - simple web application to save some basic medical records.
# Copyright (C) 2026 Jiri Skoda <developer@skodaj.cz>
#
# This program is free software; you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation; version 2 of the License.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License along
# with this program; if not, write to the Free Software Foundation, Inc.,
# 51 Franklin Street, Fifth Floor, Boston, MA 02110-1301 USA.

"""Tests for crash-safe XLSX storage."""

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.models.data import ClientCreate, HealthInsuranceCompany, HealthMeasurementCreate
from app.storage.excel import WorkbookStorage


def create_workbook(path: Path) -> None:
    """Create a minimal valid MediRec workbook."""
    workbook = Workbook()
    workbook.active.title = "HealthInsuranceCompanies"
    workbook.create_sheet("_template")
    workbook.save(path)


def test_client_round_trip(tmp_path: Path) -> None:
    """Verify client creation, measurement storage and deletion."""
    workbook_path = tmp_path / "medirec.xlsx"
    create_workbook(workbook_path)
    storage = WorkbookStorage(workbook_path)
    storage.initialize()
    storage.add_health_insurance_company(HealthInsuranceCompany(abbreviation="VZP", name="VZP", code="111"))
    client = storage.create_client(ClientCreate(surname="Škoda", name="Jiří", date_of_birth=date(1980, 1, 1), personal_id="800101/1234", sex="M", health_insurance_abbreviation="VZP"))
    storage.add_measurement(client.identifier, HealthMeasurementCreate(date=date.today(), time=time(12, 0), systolic=120, diastolic=70, pulse=70, oxygen_saturation=98, height=180, weight=80))
    assert client.sheet_name == "Skoda_Jiri_800101_1234"
    assert client.identifier
    assert len(storage.get_client(client.identifier).measurements) == 1
    assert storage.get_client(client.identifier).measurements[0].bmi_status.value == "HEALTHY"
    storage.delete_client(client.identifier)
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == ["HealthInsuranceCompanies", "_template"]
